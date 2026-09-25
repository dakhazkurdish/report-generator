import streamlit as st
import requests
import json
import io
import urllib.parse
import re
import time

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn

from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

st.set_page_config(
    page_title="سیستەمێ زیرەک یێ دروستکرنا راپورت و سمیناران", 
    page_icon="🎓", 
    layout="wide", 
    initial_sidebar_state="collapsed"
)

# ڤەشارتنا هەمی مێنیو و بارێن سێرڤەری
st.markdown("""
    <style>
    #MainMenu {visibility: hidden !important; display: none !important;}
    header {visibility: hidden !important; display: none !important;}
    footer {visibility: hidden !important; display: none !important;}
    
    .viewerBadge_container__1QSob,
    .styles_viewerBadge__1yB5G,
    [data-testid="stToolbar"],
    [data-testid="stDecoration"],
    [data-testid="stStatusWidget"],
    [data-testid="stManageAppButton"],
    .manage-app-button,
    button[title="View app in GitHub"],
    a[href*="github.com"] {
        visibility: hidden !important; 
        display: none !important;
    }
    
    .stApp { direction: rtl; text-align: right; }
    p, h1, h2, h3, label, div { text-align: right !important; font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif; }
    .stButton>button { width: 100%; border-radius: 8px; font-weight: bold; background-color: #1E3A8A; color: white; height: 3.2em; font-size: 16px; }
    textarea { direction: rtl !important; text-align: right !important; font-family: 'Segoe UI', Tahoma, sans-serif !important; }
    </style>
""", unsafe_allow_html=True)

st.title("سیستەمێ زیرەک یێ دروستکرنا راپورت و سمیناران")

raw_api_key = st.secrets.get("GEMINI_API_KEY", "")
if not raw_api_key:
    raw_api_key = st.text_input("کلیلا Gemini API لێرە بنڤیسە (دشێی چەند کلیلان ب کۆما جودا بکەی):", type="password")

with st.container():
    col1, col2 = st.columns(2)
    with col1:
        student_name = st.text_input("👤 ناڤێ قوتابی:")
        department = st.text_input("🏛️ پەیمانگەهـ یان کۆلێژ / پشک:")
        language = st.selectbox("🌐 زمانێ نڤیسینێ:", ["کوردی (بادینی)", "کوردی (سۆرانی)", "العربية", "English"])
        academic_level = st.selectbox(
            "🎓 ئاستێ ئەکادیمی یێ لێکۆڵینەوەیێ:",
            [
                "زانکۆ (بەکالۆریۆس - ستاندارد و پڕ زانیاری)",
                "پەیمانگەهـ (دبلۆم - ڕوون و کرداری)",
                "خاندنا باڵا (ماستەر و دکتۆرا - زۆر قووڵ، تیۆری و ڕەخنەگرانە)"
            ]
        )
    with col2:
        teacher_name = st.text_input("👨‍🏫 ناڤێ مامۆستایێ بابەتی:")
        topic = st.text_input("📝 بابەتێ سەرەکی یێ ڕاپۆرتێ:")
        pages_count = st.slider("📄 ژمارا لاپەڕێن پێدڤی بۆ ڕاپۆرتێ:", min_value=3, max_value=25, value=8)
        slides_count = st.slider("📊 ژمارا سلایدێن پاوەرپۆینتێ (Seminar):", min_value=5, max_value=20, value=6)

col_sub1, col_sub2 = st.columns(2)
with col_sub1:
    enable_border = st.checkbox("🖼️ چوارچێوە (Border) بۆ لاپەڕێن ڕاپۆرتا Word بهێتە دانان؟", value=True)
with col_sub2:
    uploaded_logo = st.file_uploader("🏛️ بارکرنا لۆگۆیێ فەرمی یێ زانکۆیێ (ئارەزوومەندانە):", type=["png", "jpg", "jpeg"])

custom_notes = st.text_area(
    "💡 تێبینی یان داخوازیێن تایبەت (بتنێ فەرمانە، ناچیتە ناڤ ڕاپۆرتێ):",
    placeholder="بۆ نموونە: گرنگیێ ب مێژوویا بابەتی بدە، نموونەیێن کرداری بینە، ئاستێ زانستی گەلەک بلند بیت...",
    height=80
)

def convert_numbers(text, is_rtl):
    if not is_rtl or not text:
        return text
    western_to_eastern = str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩")
    return str(text).translate(western_to_eastern)

def set_docx_rtl(paragraph, is_rtl=True):
    if not is_rtl:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        return
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    pPr.append(bidi)

def format_run(run, font_name="Calibri", size_pt=14, bold=False, color_rgb=(0, 0, 0), is_rtl=True):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color_rgb)
    if is_rtl:
        rPr = run._r.get_or_add_rPr()
        rtl = OxmlElement('w:rtl')
        rtl.set(qn('w:val'), '1')
        rPr.append(rtl)

def add_page_number_to_section(section, is_rtl):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)
    format_run(run, size_pt=10, bold=False, color_rgb=(120, 130, 140), is_rtl=is_rtl)

def add_page_borders(section):
    sectPr = section._sectPr
    pgBorders = parse_xml(r'''
        <w:pgBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:offsetFrom="page">
            <w:top w:val="single" w:sz="12" w:space="24" w:color="1E3A8A"/>
            <w:left w:val="single" w:sz="12" w:space="24" w:color="1E3A8A"/>
            <w:bottom w:val="single" w:sz="12" w:space="24" w:color="1E3A8A"/>
            <w:right w:val="single" w:sz="12" w:space="24" w:color="1E3A8A"/>
        </w:pgBorders>
    ''')
    sectPr.append(pgBorders)

def fetch_academic_logo(dept_name):
    headers = {"User-Agent": "AcademicSlideGen/8.0"}
    if dept_name:
        try:
            clean_dept = urllib.parse.quote(dept_name.strip())
            search_url = f"https://en.wikipedia.org/w/api.php?action=query&format=json&generator=search&gsrsearch={clean_dept}%20logo&gsrlimit=1&prop=pageimages&pithumbsize=400"
            r = requests.get(search_url, headers=headers, timeout=3)
            if r.status_code == 200:
                pages = r.json().get("query", {}).get("pages", {})
                for _, p_info in pages.items():
                    thumb = p_info.get("thumbnail", {}).get("source")
                    if thumb:
                        img_r = requests.get(thumb, headers=headers, timeout=3)
                        if img_r.status_code == 200 and len(img_r.content) > 1500:
                            return io.BytesIO(img_r.content)
        except Exception:
            pass
    return None

def fetch_slide_image(keyword, topic_context=""):
    headers = {"User-Agent": "AcademicSlideGen/8.0 (educational-use)"}
    term = keyword or topic_context or "academic presentation"
    try:
        clean_kw = urllib.parse.quote(term.strip())
        commons_url = f"https://commons.wikimedia.org/w/api.php?action=query&generator=search&gsrsearch={clean_kw}&gsrnamespace=6&gsrlimit=3&prop=imageinfo&iiprop=url&iiurlwidth=800&format=json"
        r = requests.get(commons_url, headers=headers, timeout=3)
        if r.status_code == 200:
            pages = r.json().get("query", {}).get("pages", {})
            for _, page in pages.items():
                title = page.get("title", "").lower()
                if any(bad in title for bad in [".svg", ".ogg", ".pdf", "flag", "icon", "logo", "map"]):
                    continue
                img_info = page.get("imageinfo", [{}])[0]
                thumb_url = img_info.get("thumburl") or img_info.get("url")
                if thumb_url:
                    img_res = requests.get(thumb_url, headers=headers, timeout=3)
                    if img_res.status_code == 200 and len(img_res.content) > 4000:
                        return io.BytesIO(img_res.content)
    except Exception:
        pass
    return None

def parse_keys(raw_input):
    return [k.strip() for k in re.split(r'[,;\s]+', raw_input) if k.strip()]

def call_gemini(prompt, keys_list):
    # مۆدێلێن کارایێن فەرمی
    candidate_models = ["gemini-2.5-flash", "gemini-2.0-flash", "gemini-1.5-flash"]
    headers = {"Content-Type": "application/json"}
    payload = {
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {
            "response_mime_type": "application/json",
            "maxOutputTokens": 8192
        }
    }
    
    last_error = ""
    for key_idx, current_key in enumerate(keys_list):
        for model_name in candidate_models:
            url = f"https://generativelanguage.googleapis.com/v1beta/models/{model_name}:generateContent?key={current_key}"
            try:
                res = requests.post(url, headers=headers, json=payload, timeout=90)
                if res.status_code == 200:
                    data = res.json()
                    candidates = data.get("candidates", [])
                    if candidates and "content" in candidates[0]:
                        parts = candidates[0]["content"].get("parts", [])
                        if parts and "text" in parts[0]:
                            text_clean = parts[0]["text"].strip()
                            text_clean = re.sub(r"^```(?:json)?\s*", "", text_clean)
                            text_clean = re.sub(r"\s*
